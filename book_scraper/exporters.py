"""Save a DataFrame as CSV, Excel, JSON or SQLite.

Clients ask for very different formats: a marketer wants Excel, a developer
wants JSON, someone tracking prices over time wants a database. Supporting
all of them is a few lines each with pandas.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATS = ("csv", "xlsx", "json", "sqlite")
EXTENSIONS = {"csv": ".csv", "xlsx": ".xlsx", "json": ".json", "sqlite": ".db"}
SQLITE_TABLE = "books"


def export(df: pd.DataFrame, fmt: str, out_dir: str | Path,
           basename: str = "books", append: bool = False) -> Path:
    """Write df in the given format and return the file path.

    append only matters for SQLite: rows are added to the existing table
    instead of replacing it, which builds a price history over several runs.
    """
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format {fmt!r}, choose from {FORMATS}")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{basename}{EXTENSIONS[fmt]}"

    if fmt == "csv":
        # utf-8-sig adds a marker so Excel opens umlauts and "£" correctly.
        df.to_csv(path, index=False, encoding="utf-8-sig")
    elif fmt == "json":
        df.to_json(path, orient="records", indent=2, force_ascii=False)
    elif fmt == "xlsx":
        _write_excel(df, path)
    elif fmt == "sqlite":
        with sqlite3.connect(path) as con:
            df.to_sql(SQLITE_TABLE, con, index=False,
                      if_exists="append" if append else "replace")
    return path


def _write_excel(df: pd.DataFrame, path: Path) -> None:
    """Excel file with a frozen header row, filters and sensible column widths."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Books")
        ws = writer.sheets["Books"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for i, column in enumerate(df.columns, start=1):
            # extra room in the header for the filter arrow
            longest = max([len(str(column)) + 3] + [len(str(v)) for v in df[column].head(200)])
            ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, 60)
        if "price" in df.columns:
            price_col = get_column_letter(df.columns.get_loc("price") + 1)
            for cell in ws[price_col][1:]:
                cell.number_format = "£#,##0.00"
