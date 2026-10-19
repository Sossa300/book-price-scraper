import json
import sqlite3

import pandas as pd
import pytest
from openpyxl import load_workbook

from book_scraper.exporters import export


@pytest.fixture
def df():
    return pd.DataFrame({"title": ["Ünïcode £ Book", "Second"], "price": [9.99, 20.0],
                         "rating": [5, 3]})


def test_csv_roundtrip(df, tmp_path):
    path = export(df, "csv", tmp_path)
    assert path.name == "books.csv"
    assert pd.read_csv(path, encoding="utf-8-sig").equals(df)


def test_json(df, tmp_path):
    data = json.loads(export(df, "json", tmp_path).read_text(encoding="utf-8"))
    assert data[0] == {"title": "Ünïcode £ Book", "price": 9.99, "rating": 5}


def test_excel_is_formatted(df, tmp_path):
    ws = load_workbook(export(df, "xlsx", tmp_path, basename="report")).active
    assert ws.title == "Books"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C3"
    assert ws["B2"].value == 9.99 and "£" in ws["B2"].number_format
    assert ws["A1"].font.bold


def test_sqlite_replace_and_append(df, tmp_path):
    path = export(df, "sqlite", tmp_path)
    export(df, "sqlite", tmp_path)  # replace: still 2 rows
    with sqlite3.connect(path) as con:
        assert con.execute("SELECT COUNT(*) FROM books").fetchone()[0] == 2
    export(df, "sqlite", tmp_path, append=True)  # append: history grows
    with sqlite3.connect(path) as con:
        assert con.execute("SELECT COUNT(*) FROM books").fetchone()[0] == 4


def test_creates_output_folder(df, tmp_path):
    path = export(df, "csv", tmp_path / "new" / "folder")
    assert path.exists()


def test_unknown_format(df, tmp_path):
    with pytest.raises(ValueError):
        export(df, "pdf", tmp_path)
